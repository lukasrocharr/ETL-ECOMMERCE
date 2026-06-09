#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Gera os entregáveis de apresentação a partir dos resultados reais da
análise Olist (10 perguntas de negócio):
  - Analise_Olist.xlsx   : dados + gráficos
  - Apresentacao_Olist.pptx : slides seguindo o Guia de Apresentação
Execução: ../.venv/bin/python gerar_relatorios.py
"""
import os
from datetime import date

# ===================== DADOS (resultados reais executados) =====================
VOLUMETRIA = [
    ("raw_orders (pedidos)", 99441),
    ("raw_order_items (itens)", 112650),
    ("raw_payments (pagamentos)", 103886),
    ("raw_reviews (avaliações)", 99224),
    ("raw_products (produtos)", 32951),
    ("raw_customers (clientes)", 99441),
    ("raw_sellers (vendedores)", 3095),
    ("raw_cat_translation", 71),
]
TOTAL_REGISTROS = sum(n for _, n in VOLUMETRIA)

PERGUNTAS = [
    "Qual forma de pagamento é mais usada?",
    "Qual a taxa de recompra dos clientes?",
    "Quais estados têm mais clientes ativos?",
    "Quais meses têm maior volume de vendas?",
    "Quais categorias foram mais vendidas?",
    "Qual o crescimento MoM de vendas?",
    "Tempo médio de entrega?",
    "Maiores motivos de cancelamento?",
    "Quais vendedores têm mais receita?",
    "Satisfação dos clientes?",
]
RESPOSTAS_CURTAS = [
    "Cartão de crédito — 73,9% das transações.",
    "Apenas 3,12% — base quase toda de compra única.",
    "SP lidera (40.302), seguido de RJ e MG.",
    "Pico em nov/2017 (Black Friday, 7.451 pedidos).",
    "cama/mesa/banho (volume) e saúde & beleza (receita).",
    "Forte alta em 2017; estabiliza em ~R$ 1M/mês em 2018.",
    "~12,5 dias; 93,2% entregues no prazo.",
    "Sem campo de 'motivo'; canceled+unavailable ≈ ruptura.",
    "Top seller (Guariba/SP): R$ 229 mil. SP domina.",
    "Nota média 4,09/5; 77% satisfeitos, 14,7% insatisfeitos.",
]

Q1 = [("credit_card", 76795, 73.92, 12542084.19),
      ("boleto", 19784, 19.04, 2869361.27),
      ("voucher", 5775, 5.56, 379436.87),
      ("debit_card", 1529, 1.47, 217989.79),
      ("not_defined", 3, 0.00, 0.00)]

Q2 = {"clientes_unicos": 96096, "clientes_recompra": 2997, "taxa": 3.12}

Q3 = [("SP", 40302, 41746), ("RJ", 12384, 12852), ("MG", 11259, 11635),
      ("RS", 5277, 5466), ("PR", 4882, 5045), ("SC", 3534, 3637),
      ("BA", 3277, 3380), ("DF", 2075, 2140), ("ES", 1964, 2033),
      ("GO", 1952, 2020), ("PE", 1609, 1652), ("CE", 1313, 1336),
      ("PA", 949, 975), ("MT", 876, 907), ("MA", 726, 747)]

Q4 = [(2017, 11, 7451, 1010271.37), (2018, 1, 7220, 950030.36),
      (2018, 3, 7188, 983213.44), (2018, 4, 6934, 996647.75),
      (2018, 5, 6853, 996517.68), (2018, 2, 6694, 844178.71),
      (2018, 8, 6452, 854686.33), (2018, 7, 6273, 895507.22),
      (2018, 6, 6160, 865124.31), (2017, 12, 5624, 743914.17),
      (2017, 10, 4568, 664219.43), (2017, 8, 4293, 573971.68)]

Q5 = [("bed_bath_table", 11115, 1036988.68), ("health_beauty", 9670, 1258681.34),
      ("sports_leisure", 8641, 988048.97), ("furniture_decor", 8334, 729762.49),
      ("computers_accessories", 7827, 911954.32), ("housewares", 6964, 632248.66),
      ("watches_gifts", 5991, 1205005.68), ("telephony", 4545, 323667.53),
      ("garden_tools", 4347, 485256.46), ("auto", 4235, 592720.11),
      ("toys", 4117, 483946.60), ("cool_stuff", 3796, 635290.85),
      ("perfumery", 3419, 399124.87), ("baby", 3065, 411764.89),
      ("electronics", 2767, 160246.74)]

Q6 = [("2016-09", 267.36, None), ("2016-10", 49507.66, 18417.23),
      ("2016-12", 10.90, -99.98), ("2017-01", 120312.87, 1103687.80),
      ("2017-02", 247303.02, 105.55), ("2017-03", 374344.30, 51.37),
      ("2017-04", 359927.23, -3.85), ("2017-05", 506071.14, 40.60),
      ("2017-06", 433038.60, -14.43), ("2017-07", 498031.48, 15.01),
      ("2017-08", 573971.68, 15.25), ("2017-09", 624401.69, 8.79),
      ("2017-10", 664219.43, 6.38), ("2017-11", 1010271.37, 52.10),
      ("2017-12", 743914.17, -26.36), ("2018-01", 950030.36, 27.71),
      ("2018-02", 844178.71, -11.14), ("2018-03", 983213.44, 16.47),
      ("2018-04", 996647.75, 1.37), ("2018-05", 996517.68, -0.01),
      ("2018-06", 865124.31, -13.19), ("2018-07", 895507.22, 3.51),
      ("2018-08", 854686.33, -4.56), ("2018-09", 145.00, -99.98)]

Q7 = {"entregues": 96470, "dias_medios": 12.5, "antes_prazo": 11.9, "pct_prazo": 93.23}

Q8a = [("delivered", 96478, 97.02), ("shipped", 1107, 1.11),
       ("canceled", 625, 0.63), ("unavailable", 609, 0.61),
       ("invoiced", 314, 0.32), ("processing", 301, 0.30),
       ("created", 5, 0.01), ("approved", 2, 0.00)]
Q8b = [("sports_leisure", 51), ("housewares", 49), ("computers_accessories", 46),
       ("health_beauty", 36), ("furniture_decor", 36), ("toys", 34),
       ("auto", 31), ("baby", 22), ("watches_gifts", 21), ("garden_tools", 19)]

Q9 = [("4869f7a5…b52b2", "SP", "Guariba", 1156, 229472.63),
      ("53243585…d8905", "BA", "Lauro de Freitas", 410, 222776.05),
      ("4a3ca931…93884", "SP", "Ibitinga", 1987, 200472.92),
      ("fa1c13f2…ecda94", "SP", "Sumaré", 586, 194042.03),
      ("7c67e144…b010ab", "SP", "Itaquaquecetuba", 1364, 187923.89),
      ("7e93a43e…bc753a", "SP", "Barueri", 340, 176431.87),
      ("da8622b1…dab84a", "SP", "Piracicaba", 1551, 160236.57),
      ("7a67c85e…03ad736", "SP", "São Paulo", 1171, 141745.53),
      ("1025f0e2…0e0bfa", "SP", "São Paulo", 1428, 138968.55),
      ("955fee92…80ce60", "SP", "São Paulo", 1499, 135171.70),
      ("46dc3b2c…d2718e", "RJ", "Rio de Janeiro", 542, 128111.19),
      ("6560211a…a7e94c0", "SP", "São Paulo", 2033, 123304.83),
      ("620c87c1…959fc6", "RJ", "Petrópolis", 798, 114774.50),
      ("7d13fca1…eb0964", "SP", "Ribeirão Preto", 578, 113628.97),
      ("5dceca12…dc4f8ca", "SP", "Santa Bárbara d'Oeste", 346, 112155.53)]

Q10a = [(5, 57328, 57.78), (4, 19142, 19.29), (3, 8179, 8.24),
        (2, 3151, 3.18), (1, 11424, 11.51)]
Q10b = {"media": 4.09, "total": 99224, "satisf": 77.07, "insatisf": 14.69}

OUT_DIR = os.path.dirname(os.path.abspath(__file__))

# ===========================================================================
#                                  EXCEL
# ===========================================================================
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.utils import get_column_letter

AZUL = "1F3864"; AZUL_CLARO = "2E5496"; CINZA = "D9E1F2"; BRANCO = "FFFFFF"
hdr_fill = PatternFill("solid", fgColor=AZUL)
sub_fill = PatternFill("solid", fgColor=AZUL_CLARO)
alt_fill = PatternFill("solid", fgColor="EEF3FB")
hdr_font = Font(bold=True, color=BRANCO, size=11)
title_font = Font(bold=True, color=AZUL, size=16)
thin = Side(style="thin", color="BBBBBB")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def estilo_cabecalho(ws, row, ncols, start_col=1, fill=hdr_fill):
    for c in range(start_col, start_col + ncols):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill; cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border


def escreve_tabela(ws, start_row, headers, rows, col_widths=None, money_cols=()):
    estilo_cabecalho(ws, start_row, len(headers))
    for j, h in enumerate(headers, start=1):
        ws.cell(row=start_row, column=j, value=h)
    for i, r in enumerate(rows, start=1):
        for j, v in enumerate(r, start=1):
            cell = ws.cell(row=start_row + i, column=j, value=v)
            cell.border = border
            if i % 2 == 0:
                cell.fill = alt_fill
            if j in money_cols and isinstance(v, (int, float)):
                cell.number_format = 'R$ #,##0.00'
            elif isinstance(v, (int, float)) and j > 1:
                cell.number_format = '#,##0' if isinstance(v, int) else '#,##0.00'
    if col_widths:
        for j, w in enumerate(col_widths, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w
    return start_row + len(rows)


wb = Workbook()

# ---- Sheet Resumo ----
ws = wb.active; ws.title = "Resumo"
ws.sheet_view.showGridLines = False
ws["A1"] = "Análise de E-commerce — Dataset Olist"; ws["A1"].font = Font(bold=True, color=AZUL, size=20)
ws["A2"] = "Projeto Final de Banco de Dados — 10 perguntas de negócio respondidas em T-SQL (SQL Server)"
ws["A2"].font = Font(italic=True, color="555555", size=11)
ws["A4"] = "Volumetria carregada"; ws["A4"].font = title_font
fim = escreve_tabela(ws, 5, ["Tabela", "Registros"], VOLUMETRIA, col_widths=[34, 16])
ws.cell(row=fim + 1, column=1, value="TOTAL").font = Font(bold=True)
ws.cell(row=fim + 1, column=2, value=TOTAL_REGISTROS).font = Font(bold=True)
ws.cell(row=fim + 1, column=2).number_format = '#,##0'

ws["D4"] = "Perguntas & Respostas (resumo)"; ws["D4"].font = title_font
escreve_tabela(ws, 5, ["#", "Pergunta", "Resposta"],
               [(i + 1, PERGUNTAS[i], RESPOSTAS_CURTAS[i]) for i in range(10)],
               col_widths=None)
ws.column_dimensions["D"].width = 5
ws.column_dimensions["E"].width = 46
ws.column_dimensions["F"].width = 52
for r in range(6, 16):
    ws.cell(row=r, column=5).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=r, column=6).alignment = Alignment(wrap_text=True, vertical="top")


def nova_aba(nome, titulo):
    s = wb.create_sheet(nome)
    s.sheet_view.showGridLines = False
    s["A1"] = titulo; s["A1"].font = title_font
    return s

# 1) Pagamentos
s = nova_aba("1_Pagamentos", "1) Formas de pagamento mais usadas")
escreve_tabela(s, 3, ["Forma", "Transações", "%", "Valor total"], Q1,
               col_widths=[16, 14, 10, 18], money_cols=(4,))
ch = PieChart(); ch.title = "Participação por forma de pagamento"
data = Reference(s, min_col=2, min_row=3, max_row=3 + len(Q1))
cats = Reference(s, min_col=1, min_row=4, max_row=3 + len(Q1))
ch.add_data(data, titles_from_data=True); ch.set_categories(cats)
ch.height = 8; ch.width = 14
s.add_chart(ch, "F3")

# 2) Recompra
s = nova_aba("2_Recompra", "2) Taxa de recompra dos clientes")
escreve_tabela(s, 3, ["Métrica", "Valor"],
               [("Clientes únicos", Q2["clientes_unicos"]),
                ("Clientes com recompra", Q2["clientes_recompra"]),
                ("Taxa de recompra (%)", Q2["taxa"])], col_widths=[28, 16])

# 3) Estados
s = nova_aba("3_Estados", "3) Estados com mais clientes ativos")
escreve_tabela(s, 3, ["Estado", "Clientes ativos", "Pedidos"], Q3, col_widths=[12, 18, 12])
ch = BarChart(); ch.type = "col"; ch.title = "Clientes ativos por estado (Top 15)"
data = Reference(s, min_col=2, min_row=3, max_row=3 + len(Q3))
cats = Reference(s, min_col=1, min_row=4, max_row=3 + len(Q3))
ch.add_data(data, titles_from_data=True); ch.set_categories(cats)
ch.height = 9; ch.width = 18; ch.legend = None
s.add_chart(ch, "F3")

# 4) Meses
s = nova_aba("4_Meses", "4) Meses com maior volume de vendas")
escreve_tabela(s, 3, ["Ano", "Mês", "Pedidos", "Receita"], Q4,
               col_widths=[8, 8, 12, 16], money_cols=(4,))

# 5) Categorias
s = nova_aba("5_Categorias", "5) Categorias mais vendidas")
escreve_tabela(s, 3, ["Categoria", "Itens vendidos", "Receita"], Q5,
               col_widths=[24, 16, 16], money_cols=(3,))
ch = BarChart(); ch.type = "bar"; ch.title = "Itens vendidos por categoria (Top 15)"
data = Reference(s, min_col=2, min_row=3, max_row=3 + len(Q5))
cats = Reference(s, min_col=1, min_row=4, max_row=3 + len(Q5))
ch.add_data(data, titles_from_data=True); ch.set_categories(cats)
ch.height = 10; ch.width = 18; ch.legend = None
s.add_chart(ch, "F3")

# 6) MoM
s = nova_aba("6_Crescimento_MoM", "6) Crescimento MoM de receita")
escreve_tabela(s, 3, ["Mês", "Receita", "Crescimento MoM (%)"], Q6,
               col_widths=[12, 16, 20], money_cols=(2,))
ch = LineChart(); ch.title = "Receita mensal (R$)"
data = Reference(s, min_col=2, min_row=3, max_row=3 + len(Q6))
cats = Reference(s, min_col=1, min_row=4, max_row=3 + len(Q6))
ch.add_data(data, titles_from_data=True); ch.set_categories(cats)
ch.height = 9; ch.width = 20; ch.legend = None
s.add_chart(ch, "F3")

# 7) Entrega
s = nova_aba("7_Entrega", "7) Tempo médio de entrega")
escreve_tabela(s, 3, ["Métrica", "Valor"],
               [("Pedidos entregues", Q7["entregues"]),
                ("Dias médios de entrega", Q7["dias_medios"]),
                ("Dias de folga vs estimativa", Q7["antes_prazo"]),
                ("% entregue no prazo", Q7["pct_prazo"])], col_widths=[30, 16])

# 8) Cancelamento
s = nova_aba("8_Cancelamento", "8) Status dos pedidos e cancelamentos")
fim = escreve_tabela(s, 3, ["Status", "Qtd", "%"], Q8a, col_widths=[16, 12, 10])
s.cell(row=fim + 3, column=1, value="Categorias mais frequentes em cancelados").font = title_font
escreve_tabela(s, fim + 4, ["Categoria", "Itens cancelados"], Q8b, col_widths=[24, 18])

# 9) Vendedores
s = nova_aba("9_Vendedores", "9) Vendedores com maior receita")
escreve_tabela(s, 3, ["Vendedor (id)", "UF", "Cidade", "Itens", "Receita"], Q9,
               col_widths=[20, 6, 22, 10, 16], money_cols=(5,))

# 10) Satisfação
s = nova_aba("10_Satisfacao", "10) Satisfação dos clientes")
escreve_tabela(s, 3, ["Nota", "Avaliações", "%"], Q10a, col_widths=[10, 14, 10])
ch = PieChart(); ch.title = "Distribuição das notas (1 a 5)"
data = Reference(s, min_col=2, min_row=3, max_row=3 + len(Q10a))
cats = Reference(s, min_col=1, min_row=4, max_row=3 + len(Q10a))
ch.add_data(data, titles_from_data=True); ch.set_categories(cats)
ch.height = 8; ch.width = 12
s.add_chart(ch, "F3")
escreve_tabela(s, 12, ["Resumo", "Valor"],
               [("Nota média", Q10b["media"]),
                ("Total de avaliações", Q10b["total"]),
                ("% satisfeitos (4-5)", Q10b["satisf"]),
                ("% insatisfeitos (1-2)", Q10b["insatisf"])], col_widths=[24, 14])

xlsx_path = os.path.join(OUT_DIR, "Analise_Olist.xlsx")
wb.save(xlsx_path)
print("OK Excel  ->", xlsx_path)

# ===========================================================================
#                                  PPTX
# ===========================================================================
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

C_AZUL = RGBColor(0x1F, 0x38, 0x64)
C_AZUL2 = RGBColor(0x2E, 0x54, 0x96)
C_ACENTO = RGBColor(0xED, 0x7D, 0x31)
C_CINZA = RGBColor(0x59, 0x59, 0x59)
C_BRANCO = RGBColor(0xFF, 0xFF, 0xFF)
C_FUNDO = RGBColor(0xF2, 0xF5, 0xFA)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def add_slide(faixa=True):
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(1, 0, 0, SW, SH)
    bg.fill.solid(); bg.fill.fore_color.rgb = C_FUNDO; bg.line.fill.background()
    bg.shadow.inherit = False
    s.shapes._spTree.remove(bg._element); s.shapes._spTree.insert(2, bg._element)
    if faixa:
        bar = s.shapes.add_shape(1, 0, 0, SW, Inches(0.18))
        bar.fill.solid(); bar.fill.fore_color.rgb = C_ACENTO; bar.line.fill.background()
        bar.shadow.inherit = False
    return s


def txt(slide, x, y, w, h, text, size=18, bold=False, color=C_CINZA,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font="Calibri"):
    tb = slide.shapes.add_textbox(x, y, w, h); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = font
    return tb


def bullets(slide, x, y, w, h, items, size=16, color=C_CINZA, space=8):
    tb = slide.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space); p.level = it[0] if isinstance(it, tuple) else 0
        texto = it[1] if isinstance(it, tuple) else it
        bold = False
        if texto.startswith("**") and texto.endswith("**"):
            texto = texto[2:-2]; bold = True
        r = p.add_run(); r.text = ("•  " if p.level == 0 else "–  ") + texto
        r.font.size = Pt(size - 2 * p.level); r.font.color.rgb = color
        r.font.bold = bold; r.font.name = "Calibri"
    return tb


def cabecalho(slide, titulo, etapa=None):
    if etapa:
        txt(slide, Inches(0.6), Inches(0.35), Inches(12), Inches(0.4), etapa,
            size=13, bold=True, color=C_ACENTO)
    txt(slide, Inches(0.6), Inches(0.7), Inches(12.1), Inches(0.9), titulo,
        size=28, bold=True, color=C_AZUL)
    ln = slide.shapes.add_shape(1, Inches(0.6), Inches(1.55), Inches(12.1), Pt(2))
    ln.fill.solid(); ln.fill.fore_color.rgb = C_AZUL2; ln.line.fill.background()
    ln.shadow.inherit = False


def kpi(slide, x, y, w, valor, rotulo):
    card = slide.shapes.add_shape(1, x, y, w, Inches(1.5))
    card.fill.solid(); card.fill.fore_color.rgb = C_BRANCO
    card.line.color.rgb = C_AZUL2; card.line.width = Pt(1); card.shadow.inherit = False
    # fonte adaptável ao comprimento do texto para não estourar o card
    n = len(valor)
    vsize = 30 if n <= 7 else (22 if n <= 11 else 16)
    txt(slide, x, y + Inches(0.18), w, Inches(0.7), valor, size=vsize, bold=True,
        color=C_AZUL, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(slide, x, y + Inches(0.95), w, Inches(0.45), rotulo, size=12,
        color=C_CINZA, align=PP_ALIGN.CENTER)


def add_chart(slide, tipo, cats, series, x, y, w, h, titulo="", legenda=False):
    cd = CategoryChartData(); cd.categories = cats
    for nome, vals in series:
        cd.add_series(nome, vals)
    gf = slide.shapes.add_chart(tipo, x, y, w, h, cd)
    ch = gf.chart
    ch.has_legend = legenda
    if legenda:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False
    if titulo:
        ch.has_title = True; ch.chart_title.text_frame.text = titulo
        ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
    try:
        plot = ch.plots[0]; plot.has_data_labels = False
    except Exception:
        pass
    return ch

# ---------- Slide 1: Capa ----------
s = add_slide(faixa=False)
faixa = s.shapes.add_shape(1, 0, 0, SW, Inches(2.6))
faixa.fill.solid(); faixa.fill.fore_color.rgb = C_AZUL; faixa.line.fill.background(); faixa.shadow.inherit = False
faixa2 = s.shapes.add_shape(1, 0, Inches(2.6), SW, Inches(0.12))
faixa2.fill.solid(); faixa2.fill.fore_color.rgb = C_ACENTO; faixa2.line.fill.background(); faixa2.shadow.inherit = False
txt(s, Inches(0.8), Inches(0.7), Inches(11.7), Inches(1.0),
    "Análise de E-commerce com Engenharia e Análise de Dados", size=34, bold=True,
    color=C_BRANCO)
txt(s, Inches(0.8), Inches(1.7), Inches(11.7), Inches(0.7),
    "Dataset Olist (Brazilian E-Commerce) — Pipeline e Análise em T-SQL / SQL Server",
    size=18, color=RGBColor(0xCF, 0xDA, 0xEC))
txt(s, Inches(0.8), Inches(3.1), Inches(11.7), Inches(0.5),
    "Projeto Final de Banco de Dados", size=20, bold=True, color=C_AZUL)
bullets(s, Inches(0.85), Inches(3.9), Inches(11.5), Inches(2.5), [
    "**Equipe (Arquiteto de Dados, Engenheiro de Dados, Analista de Dados):** [preencher integrantes]",
    "**Repositório:** github.com/lukasrocharr/ETL-ECOMMERCE",
    "**Volume analisado:** ~100 mil pedidos · 8 tabelas · {:,} registros".format(TOTAL_REGISTROS).replace(",", "."),
    "**Data:** " + date.today().strftime("%d/%m/%Y"),
], size=16)

# ---------- Slide 2: Agenda ----------
s = add_slide(); cabecalho(s, "Agenda")
bullets(s, Inches(0.9), Inches(2.0), Inches(11.5), Inches(5), [
    "**1. Introdução e Contexto** — tema, dataset e volumetria",
    "**2. Modelagem e Arquitetura** — modelo de dados, tipos e governança",
    "**3. Engenharia de Dados e ETL** — pipeline, limpeza e carga",
    "**4. Análise de Dados e Insights** — as 10 perguntas de negócio",
    "**5. Gestão, Governança e Git** — repositório e organização",
    "**6. Conclusão e Próximos Passos**",
], size=20, space=14)

# ---------- Slide 3: Introdução e Contexto ----------
s = add_slide(); cabecalho(s, "Tema, Relevância e Volumetria", "Etapa 1 — Introdução e Contexto")
bullets(s, Inches(0.7), Inches(1.9), Inches(6.6), Inches(5), [
    "**Tema:** comportamento de compra no e-commerce brasileiro.",
    "**Dataset:** Olist — pedidos reais de 2016 a 2018.",
    "**Relevância:** entender pagamento, logística, satisfação e",
    (1, "concentração geográfica para decisões de negócio."),
    "**Objetivo:** transformar dados brutos (CSV) em insights",
    (1, "acionáveis usando exclusivamente T-SQL."),
    "**Volumetria:** mais de meio milhão de registros somando",
    (1, "as 8 tabelas (geolocation adicional ~1 milhão de linhas)."),
], size=16)
# tabela volumetria
rows = VOLUMETRIA[:6]
tb = s.shapes.add_table(len(rows) + 1, 2, Inches(7.7), Inches(1.95), Inches(5.0), Inches(3.4)).table
tb.cell(0, 0).text = "Tabela"; tb.cell(0, 1).text = "Registros"
for i, (nome, n) in enumerate(rows, start=1):
    tb.cell(i, 0).text = nome; tb.cell(i, 1).text = f"{n:,}".replace(",", ".")
for j in range(2):
    cell = tb.cell(0, j); cell.fill.solid(); cell.fill.fore_color.rgb = C_AZUL
    cell.text_frame.paragraphs[0].runs[0].font.color.rgb = C_BRANCO
    cell.text_frame.paragraphs[0].runs[0].font.bold = True

# ---------- Slide 4: Perguntas de negócio ----------
s = add_slide(); cabecalho(s, "As 10 Perguntas de Negócio", "Etapa 1 — Introdução e Contexto")
col1 = [f"{i+1}. {PERGUNTAS[i]}" for i in range(5)]
col2 = [f"{i+1}. {PERGUNTAS[i]}" for i in range(5, 10)]
bullets(s, Inches(0.7), Inches(2.0), Inches(6.2), Inches(5), col1, size=16, space=14)
bullets(s, Inches(7.0), Inches(2.0), Inches(6.0), Inches(5), col2, size=16, space=14)

# ---------- Slide 5: Modelagem e Arquitetura ----------
s = add_slide(); cabecalho(s, "Modelo de Dados e Decisões", "Etapa 2 — Modelagem e Arquitetura")
bullets(s, Inches(0.7), Inches(1.9), Inches(6.7), Inches(5), [
    "**Entidades centrais:** orders (pedidos) e order_items (itens).",
    "**Dimensões:** products, sellers, customers, category_translation.",
    "**Eventos/transações:** payments, reviews.",
    "**Tipos & integridade:** carga em staging textual e conversão",
    (1, "segura com TRY_CONVERT (datas, decimais e inteiros)."),
    "**Índices/chaves:** order_id e product_id conectam as tabelas;",
    (1, "joins entre até 4 tabelas nas consultas analíticas."),
    "**Governança:** ambiente SQL Server 2022 isolado em Docker;",
    (1, "scripts versionados e reexecutáveis (idempotentes)."),
], size=15)
# mini esquema
box = s.shapes.add_shape(1, Inches(7.8), Inches(2.0), Inches(4.8), Inches(4.2))
box.fill.solid(); box.fill.fore_color.rgb = C_BRANCO; box.line.color.rgb = C_AZUL2; box.shadow.inherit = False
txt(s, Inches(7.8), Inches(2.1), Inches(4.8), Inches(0.5), "Relacionamentos (simplificado)",
    size=13, bold=True, color=C_AZUL, align=PP_ALIGN.CENTER)
bullets(s, Inches(8.05), Inches(2.7), Inches(4.4), Inches(3.4), [
    "customers → orders  (1:N)",
    "orders → order_items (1:N)",
    "order_items → products (N:1)",
    "order_items → sellers (N:1)",
    "orders → payments (1:N)",
    "orders → reviews (1:1)",
    "products → category_translation (N:1)",
], size=14, space=10)

# ---------- Slide 6: ETL ----------
s = add_slide(); cabecalho(s, "Pipeline de ETL e Tratamento de Dados", "Etapa 3 — Engenharia de Dados e ETL")
# fluxo
etapas = ["CSV brutos\n(Kaggle/Olist)", "BULK INSERT\n(staging raw_*)", "Limpeza &\nTipagem", "Consultas\nAnalíticas (DQL)"]
x = Inches(0.7); y = Inches(2.0); w = Inches(2.7); h = Inches(1.2); gap = Inches(0.45)
for i, e in enumerate(etapas):
    cx = Emu(int(x) + i * (int(w) + int(gap)))
    card = s.shapes.add_shape(1, cx, y, w, h)
    card.fill.solid(); card.fill.fore_color.rgb = C_AZUL2 if i % 2 == 0 else C_AZUL
    card.line.fill.background(); card.shadow.inherit = False
    tf = card.text_frame; tf.word_wrap = True; p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = e; r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = C_BRANCO
    if i < len(etapas) - 1:
        txt(s, Emu(int(cx) + int(w)), y + Inches(0.35), gap, Inches(0.6), "➜",
            size=24, bold=True, color=C_ACENTO, align=PP_ALIGN.CENTER)
bullets(s, Inches(0.7), Inches(3.7), Inches(12), Inches(3), [
    "**Extração robusta:** BULK INSERT com FORMAT='CSV' (trata campos entre aspas).",
    "**Valores nulos:** datas de entrega ausentes tratadas; conversão tolerante com TRY_CONVERT.",
    "**Texto sujo:** avaliações com texto livre/quebras de linha → arquivo reviews_limpo pré-tratado.",
    "**Formatos inconsistentes:** timestamps e decimais convertidos no momento da análise.",
    "**Reexecução:** scripts idempotentes (DROP IF EXISTS + recarga) garantem carga repetível.",
], size=15, space=9)

# ---------- Slides 7-11: Análise e Insights ----------
def slide_insight(etapa_titulo, kpis, chart_cfg, insight):
    s = add_slide(); cabecalho(s, etapa_titulo, "Etapa 4 — Análise de Dados e Insights")
    # KPIs
    n = len(kpis); kw = Inches(2.9)
    for i, (val, rot) in enumerate(kpis):
        kpi(s, Inches(0.7 + i * 3.05), Inches(1.85), kw, val, rot)
    # chart
    if chart_cfg:
        tipo, cats, series, titulo, legenda = chart_cfg
        add_chart(s, tipo, cats, series, Inches(0.7), Inches(3.6), Inches(6.6), Inches(3.4),
                  titulo, legenda)
        bullets(s, Inches(7.6), Inches(3.7), Inches(5.2), Inches(3.2), insight, size=15, space=10)
    else:
        bullets(s, Inches(0.7), Inches(3.7), Inches(12), Inches(3.2), insight, size=16, space=10)
    return s

# Q1 + Q2
slide_insight("Pagamentos e Recompra (P1 e P2)",
    [("73,9%", "no cartão de crédito"), ("3,12%", "taxa de recompra"),
     ("R$ 12,5 mi", "via crédito")],
    (XL_CHART_TYPE.PIE, [r[0] for r in Q1], [("Transações", [r[1] for r in Q1])],
     "Formas de pagamento", True),
    ["**Insight:** o crédito domina (73,9%) — parcelamento é chave;",
     (1, "boleto ainda relevante (19%) para o público sem cartão."),
     "**Risco:** recompra de apenas 3,12% → base de compra única.",
     "**Ação:** CRM, pós-venda e fidelização para elevar o LTV."])

# Q3 + Q4
slide_insight("Geografia e Sazonalidade (P3 e P4)",
    [("SP", "40.302 clientes"), ("Nov/2017", "pico de pedidos"),
     ("Sudeste", "concentra a demanda")],
    (XL_CHART_TYPE.COLUMN_CLUSTERED, [r[0] for r in Q3[:10]],
     [("Clientes ativos", [r[1] for r in Q3[:10]])], "Clientes ativos por estado (Top 10)", False),
    ["**Insight:** forte concentração no Sudeste (SP/RJ/MG).",
     "**Sazonalidade:** pico em nov/2017 (Black Friday).",
     "**Ação:** logística e estoque reforçados no Sudeste e",
     (1, "campanhas planejadas para datas de pico.")])

# Q5 + Q6
slide_insight("Categorias e Crescimento (P5 e P6)",
    [("cama/mesa/banho", "+ vendida"), ("health_beauty", "+ receita"),
     ("+52%", "MoM em nov/2017")],
    (XL_CHART_TYPE.LINE, [r[0] for r in Q6[3:]],
     [("Receita (R$)", [r[1] for r in Q6[3:]])], "Receita mensal", False),
    ["**Insight:** volume vem de cama/mesa/banho; receita de",
     (1, "saúde & beleza e relógios/presentes (ticket maior)."),
     "**Tendência:** crescimento forte em 2017, estabiliza em 2018.",
     "**Ação:** mix de produtos por margem e curva de demanda."])

# Q7 + Q8
slide_insight("Logística e Cancelamentos (P7 e P8)",
    [("12,5 dias", "entrega média"), ("93,2%", "no prazo"),
     ("0,6%", "cancelados")],
    (XL_CHART_TYPE.COLUMN_CLUSTERED, [r[0] for r in Q8a[:5]],
     [("Pedidos", [r[1] for r in Q8a[:5]])], "Status dos pedidos (Top 5)", False),
    ["**Insight:** logística saudável — 93% no prazo, ~12 dias.",
     "**Cancelamento:** o Olist não traz 'motivo' textual;",
     (1, "'unavailable' (0,6%) sugere ruptura de estoque."),
     "**Ação:** monitorar estoque e prazos das categorias críticas."])

# Q9 + Q10
slide_insight("Vendedores e Satisfação (P9 e P10)",
    [("R$ 229 mil", "top seller"), ("4,09/5", "nota média"),
     ("14,7%", "insatisfeitos")],
    (XL_CHART_TYPE.PIE, [str(r[0]) for r in Q10a],
     [("Avaliações", [r[1] for r in Q10a])], "Notas de avaliação (1-5)", True),
    ["**Insight:** receita concentrada em poucos sellers de SP.",
     "**Satisfação:** boa média (4,09), mas 14,7% dão nota 1-2.",
     (1, "11,5% dão nota 1 — cauda de insatisfação relevante."),
     "**Ação:** plano de qualidade para sellers e SAC proativo."])

# ---------- Slide 12: Gestão e Git ----------
s = add_slide(); cabecalho(s, "Organização, Governança e Git", "Etapa 5 — Gestão e Governança")
bullets(s, Inches(0.7), Inches(1.9), Inches(12), Inches(5), [
    "**Repositório:** github.com/lukasrocharr/ETL-ECOMMERCE (público, com README).",
    "**Estrutura de pastas:** sql/ (DDL, ETL, objetos), analise_olist/ (carga + 10 perguntas), docs/.",
    "**Scripts:** 00_setup_e_carga.sql (carga) e 01_perguntas.sql (as 10 análises).",
    "**Reprodutibilidade:** ambiente em Docker + scripts idempotentes = qualquer membro reproduz.",
    "**Versionamento:** histórico de commits documentando a evolução do projeto.",
    "**Gestão:** quadro Kanban (GitHub Projects/Trello) — [inserir print do board].",
], size=17, space=11)

# ---------- Slide 13: Conclusão ----------
s = add_slide(); cabecalho(s, "Conclusão e Próximos Passos", "Etapa 6 — Encerramento")
bullets(s, Inches(0.7), Inches(1.9), Inches(6.7), Inches(5), [
    "**Principais conclusões:**",
    (1, "Crédito domina; logística forte (93% no prazo)."),
    (1, "Demanda concentrada no Sudeste e em datas de pico."),
    (1, "Recompra baixíssima (3,1%) e 14,7% de insatisfeitos."),
], size=16)
bullets(s, Inches(7.4), Inches(1.9), Inches(5.4), Inches(5), [
    "**Próximos passos:**",
    (1, "Modelar um Data Warehouse estrela (fato vendas)."),
    (1, "Conectar Dashboard de BI (Power BI/Metabase)."),
    (1, "Programa de fidelização para elevar a recompra."),
    (1, "Painel de qualidade de sellers e satisfação."),
], size=16)

# ---------- Slide 14: Obrigado ----------
s = add_slide(faixa=False)
faixa = s.shapes.add_shape(1, 0, 0, SW, SH)
faixa.fill.solid(); faixa.fill.fore_color.rgb = C_AZUL; faixa.line.fill.background(); faixa.shadow.inherit = False
txt(s, Inches(1), Inches(2.7), Inches(11.3), Inches(1.2), "Obrigado!", size=48, bold=True,
    color=C_BRANCO, align=PP_ALIGN.CENTER)
txt(s, Inches(1), Inches(4.0), Inches(11.3), Inches(0.8),
    "Perguntas?  ·  github.com/lukasrocharr/ETL-ECOMMERCE", size=20,
    color=RGBColor(0xCF, 0xDA, 0xEC), align=PP_ALIGN.CENTER)

pptx_path = os.path.join(OUT_DIR, "Apresentacao_Olist.pptx")
prs.save(pptx_path)
print("OK PPTX   ->", pptx_path)
print("Slides:", len(prs.slides._sldIdLst))
